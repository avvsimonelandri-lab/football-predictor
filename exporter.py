"""Esportazione dei pronostici della settimana in Excel e PDF."""
from __future__ import annotations

import os
from datetime import date, datetime, timedelta

from predictor import FootballPredictor, best_value

GIORNI = ["Lun", "Mar", "Mer", "Gio", "Ven", "Sab", "Dom"]


def collect(data: dict, days: int = 7, leagues: list | None = None) -> list[dict]:
    """Pronostici di tutte le partite dei prossimi `days` giorni."""
    blend_w = ((data.get("blend") or {}).get("w"))
    xg_beta = ((data.get("xg") or {}).get("beta"))
    today = date.today().isoformat()
    limit = (date.today() + timedelta(days=days)).isoformat()
    rows = []
    for name, ld in data.get("leagues", {}).items():
        if leagues and name not in leagues:
            continue
        fx = [f for f in ld.get("fixtures", []) if today <= f["date"] <= limit]
        if not fx:
            continue
        m = FootballPredictor(ld, blend_w=blend_w, xg_beta=xg_beta)
        for f in fx:
            r = m.predict_detailed(f["home"], f["away"], f.get("odds"))
            if not r:
                continue
            p = m.best_1x2(r)
            over = ((r.get("blend") or {}).get("UnderOver25") or r["UnderOver25"])["Over"]
            (sh, sa), sp = r["exact_scores"][0]
            bv = best_value(r)
            ang = (r.get("Stats") or {}).get("Angoli")
            car = (r.get("Stats") or {}).get("Cartellini")
            pt = r.get("PrimoTempo") or {}
            o = f.get("odds") or {}
            d = date.fromisoformat(f["date"])
            rows.append({
                "league": name, "date": f["date"], "day": f"{GIORNI[d.weekday()]} {d:%d/%m}", "time": f.get("time", ""),
                "home": f["home"], "away": f["away"],
                "p1": p["1"], "px": p["X"], "p2": p["2"],
                "pron": max(p, key=p.get), "conf": max(p.values()),
                "dc": max(r["DoppiaChance"], key=r["DoppiaChance"].get),
                "over": over, "gg": r["GoalNoGoal"]["Goal"], "score": f"{sh}-{sa}", "score_p": sp,
                "xg_h": r["xg"][0], "xg_a": r["xg"][1],
                "pt1x2": max(pt.get("1X2", {"-": 0}), key=pt.get("1X2", {"-": 0}).get) if pt else "",
                "ang": sum(ang["exp"]) if ang else None, "car": sum(car["exp"]) if car else None,
                "q1": o.get("H"), "qx": o.get("D"), "q2": o.get("A"),
                "value": f"{bv['market']} @{bv['odds']:.2f} (+{bv['ev'] * 100:.0f}%)" if bv else "",
            })
    rows.sort(key=lambda r: (r["date"], r["time"], r["league"]))
    return rows


# ---------------------------------------------------------------------- Excel
def to_excel(rows: list, path: str):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Pronostici"
    head = ["Campionato", "Giorno", "Ora", "Casa", "Ospite", "1", "X", "2", "Pron.", "Doppia ch.", "Over 2.5",
            "Goal", "Ris. più prob.", "Gol attesi casa", "Gol attesi ospite", "1X2 1°T", "Angoli attesi",
            "Cartellini attesi", "Quota 1", "Quota X", "Quota 2", "Valore"]
    ws.append([f"Pronostici Football Predictor Pro — creato il {datetime.now():%d/%m/%Y %H:%M}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(head)
    hdr_fill = PatternFill("solid", fgColor="1F2937")
    for c in ws[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    gold = PatternFill("solid", fgColor="FFF3C4")
    for r in rows:
        ws.append([r["league"], r["day"], r["time"], r["home"], r["away"], r["p1"], r["px"], r["p2"], r["pron"],
                   r["dc"], r["over"], r["gg"], r["score"], round(r["xg_h"], 2), round(r["xg_a"], 2), r["pt1x2"],
                   round(r["ang"], 1) if r["ang"] else None, round(r["car"], 1) if r["car"] else None,
                   r["q1"], r["qx"], r["q2"], r["value"]])
        row = ws.max_row
        for col in (6, 7, 8, 11, 12):
            ws.cell(row, col).number_format = "0.0%"
        best_col = {"1": 6, "X": 7, "2": 8}[r["pron"]]
        ws.cell(row, best_col).font = Font(bold=True)
        if r["value"]:
            ws.cell(row, 22).fill = gold
    widths = [18, 10, 6, 22, 22, 7, 7, 7, 6, 9, 8, 7, 10, 9, 9, 7, 9, 9, 7, 7, 7, 22]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "F4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(head))}{max(3, ws.max_row)}"
    ws.row_dimensions[3].height = 32

    vs = wb.create_sheet("Segnalazioni valore")
    vs.append(["Campionato", "Giorno", "Ora", "Partita", "Segnalazione"])
    for c in vs[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = hdr_fill
    for r in rows:
        if r["value"]:
            vs.append([r["league"], r["day"], r["time"], f"{r['home']} - {r['away']}", r["value"]])
    for i, w in enumerate([18, 10, 6, 40, 26], 1):
        vs.column_dimensions[get_column_letter(i)].width = w
    vs.append([])
    vs.append(["Stime statistiche, non certezze. Il gioco è vietato ai minori e può causare dipendenza."])
    wb.save(path)


# ---------------------------------------------------------------------- PDF
def _font_paths():
    cands = [
        (r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\arialbd.ttf"),
        ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ]
    for reg, bold in cands:
        if os.path.exists(reg) and os.path.exists(bold):
            return reg, bold
    return None


def to_pdf(rows: list, path: str):
    from fpdf import FPDF

    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(True, margin=12)
    fonts = _font_paths()
    if fonts:
        pdf.add_font("F", "", fonts[0])
        pdf.add_font("F", "B", fonts[1])
        fam = "F"
    else:
        fam = "Helvetica"

    def txt(s):
        s = str(s)
        return s if fonts else s.encode("latin-1", "replace").decode("latin-1")

    pdf.add_page()
    pdf.set_font(fam, "B", 15)
    pdf.cell(0, 8, txt("Football Predictor Pro — pronostici della settimana"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font(fam, "", 9)
    pdf.set_text_color(110)
    pdf.cell(0, 5, txt(f"Creato il {datetime.now():%d/%m/%Y %H:%M}. Dove ci sono le quote, 1X2 e Over sono la stima combinata modello + bookmaker."),
             new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(0)
    cols = [("Giorno", 20), ("Ora", 11), ("Partita", 70), ("1", 13), ("X", 13), ("2", 13), ("Pron.", 12),
            ("DC", 11), ("Over 2.5", 16), ("Goal", 13), ("Ris.", 12), ("Angoli", 14), ("Quote 1-X-2", 30), ("Valore", 0)]
    league = None
    for r in sorted(rows, key=lambda r: (r["league"], r["date"], r["time"])):
        if r["league"] != league:
            league = r["league"]
            pdf.ln(3)
            pdf.set_font(fam, "B", 11)
            pdf.set_fill_color(31, 41, 55)
            pdf.set_text_color(255)
            pdf.cell(0, 7, txt(league), fill=True, new_x="LMARGIN", new_y="NEXT")
            pdf.set_text_color(90)
            pdf.set_font(fam, "B", 8)
            for name, w in cols:
                pdf.cell(w, 5, txt(name), border="B")
            pdf.ln()
            pdf.set_text_color(0)
        pdf.set_font(fam, "", 8.5)
        q = f"{r['q1']:.2f} {r['qx']:.2f} {r['q2']:.2f}" if r["q1"] else "-"
        vals = [r["day"], r["time"], f"{r['home']} - {r['away']}", f"{r['p1'] * 100:.0f}%", f"{r['px'] * 100:.0f}%",
                f"{r['p2'] * 100:.0f}%", r["pron"], r["dc"], f"{r['over'] * 100:.0f}%", f"{r['gg'] * 100:.0f}%",
                r["score"], f"{r['ang']:.1f}" if r["ang"] else "-", q, r["value"]]
        for (name, w), v in zip(cols, vals):
            bold = (name == {"1": "1", "X": "X", "2": "2"}[r["pron"]]) or (name == "Valore" and v)
            pdf.set_font(fam, "B" if bold else "", 8.5)
            pdf.cell(w, 5.5, txt(v)[:48])
        pdf.ln()
    pdf.ln(4)
    pdf.set_font(fam, "", 7.5)
    pdf.set_text_color(110)
    pdf.multi_cell(0, 4, txt("Pron. = esito più probabile; DC = doppia chance più probabile; Ris. = risultato esatto più "
                             "probabile; Valore = la quota paga più di quanto la stima ritiene giusto. Stime statistiche, "
                             "non certezze: il modello non conosce infortuni e formazioni. Il gioco è vietato ai minori e "
                             "può causare dipendenza."))
    pdf.output(path)


def open_file(path: str):
    try:
        if os.name == "nt":
            os.startfile(path)  # type: ignore[attr-defined]
    except Exception:
        pass
